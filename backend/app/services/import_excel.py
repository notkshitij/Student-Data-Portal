"""
Service for importing Excel data into the Student Data Verification Portal.

Excel format rules:
  - First row = headers (column names). Headers must never contain [COLLECT].
  - ``Email ID*`` is the required portal login identity column.
  - Subsequent rows are student records.
  - A data cell whose stripped, case-insensitive value is exactly ``[COLLECT]``
    means the student must provide that value later.  The marker is NOT stored
    as imported data.
  - Ordinary text that merely *contains* ``[COLLECT]`` (e.g. "My [COLLECT]
    information") is treated as normal imported data.
  - Different students may have [COLLECT] in different columns.

Formula handling:
  openpyxl is opened with ``data_only=True`` so that cached formula results
  are used.  If a formula cell has no cached value, ``None`` is stored.

Blank cells that are not [COLLECT] are stored as ``None`` (no collection
requirement is inferred from emptiness).
"""

import io
import re
import uuid
from datetime import datetime, timezone

from openpyxl import load_workbook
from pydantic import TypeAdapter, EmailStr, ValidationError
from sqlalchemy.orm import Session

from app.models.audit_log import AuditLog
from app.models.campaign import Campaign, CampaignStatus
from app.models.campaign_field import CampaignField
from app.models.campaign_student import CampaignStudent
from app.models.import_record import Import, ImportStatus
from app.models.imported_field_value import ImportedFieldValue
from app.models.user import User, UserRole

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

IDENTITY_COLUMN = "Email ID*"
"""The exact header name (case-insensitive match) used as the portal login
identity column."""

ALLOWED_DOMAINS = frozenset({"poornima.edu.in", "poornima.org"})

# Configurable limits to prevent resource exhaustion
MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_ROWS = 10_000
MAX_COLUMNS = 200

_email_adapter = TypeAdapter(EmailStr)

# Pre-compiled pattern: matches *exactly* [COLLECT] after stripping whitespace
_COLLECT_RE = re.compile(r"^\s*\[COLLECT\]\s*$", re.IGNORECASE)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def process_excel_import(
    db: Session,
    file_bytes: bytes,
    original_filename: str,
    stored_filename: str,
    campaign: Campaign,
    admin_id: uuid.UUID,
) -> dict:
    """Process an Excel workbook and populate a campaign with student records.

    The entire operation is expected to run inside a transaction managed by the
    caller.  On validation failure a ``ValueError`` is raised so the caller can
    roll back.

    Returns a summary dict suitable for the API response.
    """

    # --- File-size guard ------------------------------------------------
    if len(file_bytes) > MAX_FILE_BYTES:
        raise ValueError(
            f"File exceeds the maximum allowed size of "
            f"{MAX_FILE_BYTES // (1024 * 1024)} MB"
        )

    # --- Open workbook --------------------------------------------------
    try:
        wb = load_workbook(
            filename=io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
        sheet = wb.active
    except Exception:
        raise ValueError("The uploaded file is not a valid Excel workbook")

    # --- Read and validate headers --------------------------------------
    rows_iter = sheet.iter_rows()
    try:
        header_cells = next(rows_iter)
    except StopIteration:
        raise ValueError("The workbook contains no data")

    headers = _extract_headers(header_cells)

    # --- Locate the identity column -------------------------------------
    email_col_idx = _find_identity_column(headers)

    # --- Read all data rows upfront for validation ----------------------
    data_rows = _read_data_rows(rows_iter, len(headers))
    wb.close()

    if not data_rows:
        raise ValueError("The workbook contains no student rows after the header")

    # --- Validate every row *before* touching the database --------------
    validated_rows = _validate_rows(data_rows, email_col_idx, headers)

    # --- Persist ---------------------------------------------------------
    return _persist_import(
        db=db,
        headers=headers,
        validated_rows=validated_rows,
        email_col_idx=email_col_idx,
        campaign=campaign,
        original_filename=original_filename,
        stored_filename=stored_filename,
        admin_id=admin_id,
    )


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------


def _extract_headers(header_cells) -> list[str]:
    """Extract and validate header names from the first row.

    Validation rules:
      - Every header cell must be non-blank.
      - ``[COLLECT]`` is not allowed in any header — neither as the
        exact value nor as a substring (case-insensitive, after trimming).
      - Duplicate headers (compared case-insensitively after trimming)
        cause the import to be rejected.
    """
    headers: list[str] = []
    for cell in header_cells:
        val = cell.value
        if val is None or str(val).strip() == "":
            raise ValueError(
                "Header row contains blank cells — every column must have "
                "a name"
            )
        name = str(val).strip()

        # [COLLECT] must never appear in a header — not as the exact
        # value and not as a substring.
        if "[collect]" in name.lower():
            raise ValueError(
                f"[COLLECT] is not allowed in column headers. "
                f"Found in header: '{name}'"
            )

        headers.append(name)

    if not headers:
        raise ValueError("No headers found in the first row")

    if len(headers) > MAX_COLUMNS:
        raise ValueError(
            f"Too many columns ({len(headers)}); maximum is {MAX_COLUMNS}"
        )

    # Duplicate header check (case-insensitive, trimmed)
    seen: dict[str, int] = {}
    for idx, h in enumerate(headers):
        key = h.lower()
        if key in seen:
            raise ValueError(
                f"Duplicate column header detected: '{h}' appears in "
                f"columns {seen[key] + 1} and {idx + 1}"
            )
        seen[key] = idx

    return headers


def _find_identity_column(headers: list[str]) -> int:
    """Return the 0-based index of the ``Email ID*`` column."""
    target = IDENTITY_COLUMN.lower()
    for i, h in enumerate(headers):
        if h.lower() == target:
            return i
    raise ValueError(
        f"Required identity column '{IDENTITY_COLUMN}' not found in headers"
    )


def _read_data_rows(rows_iter, num_headers: int) -> list[list]:
    """Read all data rows, enforcing the row limit."""
    data_rows: list[list] = []
    for row_cells in rows_iter:
        values = [cell.value for cell in row_cells]
        # Pad or truncate to match header count
        if len(values) < num_headers:
            values.extend([None] * (num_headers - len(values)))
        elif len(values) > num_headers:
            values = values[:num_headers]

        # Skip completely empty rows
        if all(v is None or str(v).strip() == "" for v in values):
            continue

        data_rows.append(values)

        if len(data_rows) > MAX_ROWS:
            raise ValueError(
                f"Workbook exceeds the maximum of {MAX_ROWS} student rows"
            )

    return data_rows


def _cell_to_string(value) -> str | None:
    """Convert a cell value to a string, preserving useful types.

    - ``None`` → ``None``
    - ``datetime`` → ISO 8601 string
    - ``int`` / ``float`` → string (integers render without decimals)
    - Everything else → ``str(value)``
    """
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.isoformat()

    if isinstance(value, float):
        # Render 1234.0 as "1234" to avoid misleading decimal noise
        if value == int(value):
            return str(int(value))
        return str(value)

    return str(value)


def _is_collect_marker(value) -> bool:
    """Return True only if the value is exactly the [COLLECT] marker."""
    if value is None:
        return False
    return bool(_COLLECT_RE.match(str(value)))


def _validate_rows(
    data_rows: list[list],
    email_col_idx: int,
    headers: list[str],
) -> list[dict]:
    """Validate every data row and return structured validated data.

    Raises ``ValueError`` with a descriptive message on the first fatal error.
    Returns a list of dicts with keys: ``email``, ``cell_values``.
    """
    seen_emails: dict[str, int] = {}  # email -> 1-based row number
    validated: list[dict] = []

    for row_idx, row_values in enumerate(data_rows):
        display_row = row_idx + 2  # 1-based, accounting for header row

        raw_email = row_values[email_col_idx]

        # --- Email ID* must not be blank --------------------------------
        if raw_email is None or str(raw_email).strip() == "":
            raise ValueError(
                f"Row {display_row}: '{IDENTITY_COLUMN}' is empty"
            )

        # --- Email ID* must not be [COLLECT] ----------------------------
        if _is_collect_marker(raw_email):
            raise ValueError(
                f"Row {display_row}: '{IDENTITY_COLUMN}' cannot be [COLLECT]"
            )

        # --- Normalize and validate the email address -------------------
        email_str = str(raw_email).strip().lower()
        try:
            email_str = _email_adapter.validate_python(email_str)
        except ValidationError:
            raise ValueError(
                f"Row {display_row}: '{IDENTITY_COLUMN}' contains an invalid "
                f"email address"
            )

        # --- Domain check -----------------------------------------------
        domain = email_str.rsplit("@", 1)[-1]
        if domain not in ALLOWED_DOMAINS:
            raise ValueError(
                f"Row {display_row}: email domain '{domain}' is not allowed. "
                f"Only {', '.join(sorted(ALLOWED_DOMAINS))} are accepted"
            )

        # --- Duplicate email within the same workbook -------------------
        if email_str in seen_emails:
            raise ValueError(
                f"Row {display_row}: duplicate '{IDENTITY_COLUMN}' — "
                f"'{email_str}' already appears in row {seen_emails[email_str]}"
            )
        seen_emails[email_str] = display_row

        # --- Prepare cell values ----------------------------------------
        cell_values: list[dict] = []
        for col_idx, value in enumerate(row_values):
            is_collect = _is_collect_marker(value)
            cell_values.append({
                "value": None if is_collect else _cell_to_string(value),
                "requires_input": is_collect,
            })

        validated.append({
            "email": email_str,
            "cell_values": cell_values,
        })

    return validated


def _persist_import(
    db: Session,
    headers: list[str],
    validated_rows: list[dict],
    email_col_idx: int,
    campaign: Campaign,
    original_filename: str,
    stored_filename: str,
    admin_id: uuid.UUID,
) -> dict:
    """Write validated data to the database.

    Assumes the caller wraps this in a transaction / savepoint.
    """
    # --- Clean up existing import data for a clean replacement -----------
    # Since we are replacing, we must delete old fields, students, and imports
    # The cascading deletes will handle ImportedFieldValues and StudentResponses
    if campaign.fields:
        db.query(CampaignField).filter(CampaignField.campaign_id == campaign.id).delete()
    if campaign.students:
        db.query(CampaignStudent).filter(CampaignStudent.campaign_id == campaign.id).delete()
    if campaign.imports:
        db.query(Import).filter(Import.campaign_id == campaign.id).delete()
    db.flush()

    # --- Import record --------------------------------------------------
    import_record = Import(
        campaign_id=campaign.id,
        uploaded_by_id=admin_id,
        original_filename=original_filename,
        stored_filename=stored_filename,
        status=ImportStatus.PROCESSING,
    )
    db.add(import_record)
    db.flush()

    # --- Campaign fields (preserve column order) ------------------------
    campaign_fields: list[CampaignField] = []
    for order, header in enumerate(headers):
        field = CampaignField(
            campaign_id=campaign.id,
            field_name=header,
            field_order=order,
        )
        db.add(field)
        campaign_fields.append(field)
    db.flush()

    # --- Students and imported values -----------------------------------
    num_students = 0
    num_collect_cells = 0
    fields_with_collect: set[str] = set()

    for row_data in validated_rows:
        email = row_data["email"]

        # Get or create User
        user = db.query(User).filter(User.email == email).first()
        if not user:
            user = User(
                email=email,
                role=UserRole.STUDENT,
                is_active=True,
                # google_subject_id intentionally left NULL
            )
            db.add(user)
            db.flush()

        # Campaign membership (fail on duplicates within the same campaign)
        existing_membership = (
            db.query(CampaignStudent)
            .filter_by(campaign_id=campaign.id, student_id=user.id)
            .first()
        )
        if existing_membership:
            raise ValueError(
                f"Duplicate campaign membership for email '{email}'"
            )

        cs = CampaignStudent(
            campaign_id=campaign.id,
            student_id=user.id,
        )
        db.add(cs)
        db.flush()
        num_students += 1

        # Imported field values
        for field, cell_data in zip(campaign_fields, row_data["cell_values"]):
            ifv = ImportedFieldValue(
                campaign_student_id=cs.id,
                campaign_field_id=field.id,
                imported_value=cell_data["value"],
                requires_student_input=cell_data["requires_input"],
            )
            db.add(ifv)

            if cell_data["requires_input"]:
                num_collect_cells += 1
                fields_with_collect.add(field.field_name)

    # --- Finalize import record -----------------------------------------
    import_record.status = ImportStatus.COMPLETED
    import_record.row_count = num_students
    import_record.completed_at = datetime.now(timezone.utc)

    # --- Audit log ------------------------------------------------------
    audit = AuditLog(
        user_id=admin_id,
        action="IMPORT_CAMPAIGN",
        entity_type="Campaign",
        entity_id=str(campaign.id),
        details={
            "filename": original_filename,
            "students_imported": num_students,
            "fields_count": len(campaign_fields),
            "collect_cells": num_collect_cells,
        },
    )
    db.add(audit)

    return {
        "campaign_id": campaign.id,
        "import_id": import_record.id,
        "campaign_name": campaign.name,
        "status": import_record.status.value,
        "num_students": num_students,
        "num_fields": len(campaign_fields),
        "num_collect_cells": num_collect_cells,
        "fields_discovered": headers,
        "fields_with_collect": sorted(fields_with_collect),
        "rows_processed": len(validated_rows),
        "rejected_rows": 0,
    }
